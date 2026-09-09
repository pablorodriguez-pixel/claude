#!/usr/bin/env python3
"""
Growth OS — generador del Excel a partir de los YAML del repo.

    python3 growth-os/build_excel.py [salida.xlsx]

Fuente de la verdad: growth-os/data/*.yaml
Salida por defecto:  growth-os/dist/Founderz_Growth_OS.xlsx

El Excel resultante NO contiene valores calculados a mano: todo lo derivado
son fórmulas, de modo que el fichero sigue siendo correcto cuando el equipo
edita los datos directamente en Excel.
"""
import sys
import datetime as dt
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"

# --------------------------------------------------------------------------
# Marca Founderz (ver CLAUDE.md — tokens web)
# --------------------------------------------------------------------------
PURPLE = "5045C8"
PURPLE_DARK = "2F2976"
LAVENDER = "EDEBFA"
INK = "111115"
GREY = "5A5A5A"
BORDER_C = "DCDCE6"
COMPUTED_BG = "F6F6FB"
WHITE = "FFFFFF"

GREEN_BG, GREEN_FG = "DFF0E8", "1F5D45"
AMBER_BG, AMBER_FG = "FDECD3", "8A4B00"
RED_BG, RED_FG = "FBDDDA", "8C1F16"
NEUTRAL_BG, NEUTRAL_FG = "F0F0F4", "6E6E78"
PURPLE_BG, PURPLE_FG = "E4E1F7", PURPLE_DARK

FONT = "Trebuchet MS"  # fallback oficial de RundDisplay, presente en Office

# --------------------------------------------------------------------------
# Tamaño de los bloques de datos: filas reales + holgura para que el equipo
# pueda añadir filas en Excel sin tener que regenerar el fichero.
# --------------------------------------------------------------------------
P_FIRST, P_LAST = 2, 41    # Proyectos
K_FIRST, K_LAST = 2, 71    # KPIs
H_FIRST, H_LAST = 2, 81    # Hitos

FMT_EUR = '#,##0 €'
FMT_PCT0 = '0%'
FMT_PCT1 = '0.0%'
FMT_PCT_IN = '0.#%'   # columnas de entrada: no redondea a entero al exportar
FMT_DATE = 'dd/mm/yyyy'
FMT_INT = '#,##0'
FMT_NUM = '#,##0.###'


# --------------------------------------------------------------------------
# Helpers de estilo
# --------------------------------------------------------------------------
def thin(color=BORDER_C):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def title(ws, cell, text, size=20):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT, size=size, bold=True, color=PURPLE_DARK)
    return c


def subtitle(ws, cell, text, size=10):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT, size=size, color=GREY)
    return c


def section(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text.upper())
    c.font = Font(name=FONT, size=9, bold=True, color=PURPLE)
    c.alignment = Alignment(horizontal="left", vertical="center")
    return c


def table_header(ws, row, first_col, headers, computed_from=None):
    """Escribe una fila de encabezados. computed_from = índice (1-based dentro
    de headers) a partir del cual las columnas son calculadas."""
    for i, h in enumerate(headers):
        col = first_col + i
        c = ws.cell(row=row, column=col, value=h)
        is_calc = computed_from is not None and (i + 1) >= computed_from
        c.font = Font(name=FONT, size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=PURPLE_DARK if is_calc else PURPLE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin(PURPLE_DARK)
    ws.row_dimensions[row].height = 34


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def tile(ws, row, col, label, formula, number_format=FMT_INT, note=None):
    """Tarjeta de KPI: etiqueta arriba, valor grande debajo."""
    lab = ws.cell(row=row, column=col, value=label)
    lab.font = Font(name=FONT, size=8, bold=True, color=PURPLE)
    lab.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
    lab.fill = PatternFill("solid", fgColor=LAVENDER)
    lab.border = Border(left=Side("thin", color=LAVENDER), right=Side("thin", color=LAVENDER),
                        top=Side("thin", color=LAVENDER))
    val = ws.cell(row=row + 1, column=col, value=formula)
    val.font = Font(name=FONT, size=16, bold=True, color=INK)
    val.alignment = Alignment(horizontal="left", vertical="center")
    val.number_format = number_format
    val.fill = PatternFill("solid", fgColor=LAVENDER)
    val.border = Border(left=Side("thin", color=LAVENDER), right=Side("thin", color=LAVENDER),
                        bottom=Side("thin", color=LAVENDER))
    if note:
        lab.comment = Comment(note, "Growth OS")
    ws.row_dimensions[row].height = 26
    ws.row_dimensions[row + 1].height = 30


def state_rules(ws, rng, mapping):
    """Formato condicional por valor exacto de texto."""
    for value, (bg, fg) in mapping.items():
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="equal", formula=['"%s"' % value],
                       fill=PatternFill("solid", bgColor=bg),
                       font=Font(name=FONT, size=9, bold=True, color=fg)),
        )


ESTADO_COLORS = {
    "Backlog": (NEUTRAL_BG, NEUTRAL_FG),
    "En curso": (PURPLE_BG, PURPLE_FG),
    "En riesgo": (AMBER_BG, AMBER_FG),
    "Bloqueado": (RED_BG, RED_FG),
    "Completado": (GREEN_BG, GREEN_FG),
    "Cancelado": (NEUTRAL_BG, NEUTRAL_FG),
}
SEMAFORO_COLORS = {
    "Vencido": (RED_BG, RED_FG),
    "Desviado": (AMBER_BG, AMBER_FG),
    "Inminente": (AMBER_BG, AMBER_FG),
    "En plazo": (GREEN_BG, GREEN_FG),
}
RIESGO_COLORS = {
    "Alto": (RED_BG, RED_FG),
    "Medio": (AMBER_BG, AMBER_FG),
    "Bajo": (GREEN_BG, GREEN_FG),
}
KPI_COLORS = {
    "Conseguido": (GREEN_BG, GREEN_FG),
    "En camino": (PURPLE_BG, PURPLE_FG),
    "Atención": (AMBER_BG, AMBER_FG),
    "Crítico": (RED_BG, RED_FG),
}
ALERTA_COLORS = {
    "Vencido": (RED_BG, RED_FG),
    "Esta semana": (AMBER_BG, AMBER_FG),
    "Próximos 30 días": (PURPLE_BG, PURPLE_FG),
    "En plazo": (GREEN_BG, GREEN_FG),
}


# --------------------------------------------------------------------------
# Rangos absolutos reutilizados en las fórmulas
# --------------------------------------------------------------------------
def rng(sheet, col, first, last):
    return "%s!$%s$%d:$%s$%d" % (sheet, col, first, col, last)


P = lambda col: rng("Proyectos", col, P_FIRST, P_LAST)
K = lambda col: rng("KPIs", col, K_FIRST, K_LAST)
H = lambda col: rng("Hitos", col, H_FIRST, H_LAST)

# Columnas de Proyectos
P_ID, P_NOM, P_DEP, P_SUB, P_OWN = "A", "B", "C", "D", "E"
P_TIPO, P_PRIO, P_EST, P_FASE = "F", "G", "H", "I"
P_INI, P_FIN, P_AVA, P_RIE = "J", "K", "L", "M"
P_PPTO, P_GASTO = "N", "O"
P_CONS, P_DIAS, P_AVAE, P_DESV, P_SEM = "S", "T", "U", "V", "W"
P_NKPI, P_PKPI, P_NHIT, P_HOK, P_HVEN = "X", "Y", "Z", "AA", "AB"

# Columnas de KPIs
K_ID, K_PID, K_NOM, K_UNI, K_DIR = "A", "B", "C", "D", "E"
K_BASE, K_TGT, K_ACT, K_FREQ, K_SRC = "F", "G", "H", "I", "J"
K_PROY, K_DEP, K_SUB, K_PCT, K_EST = "K", "L", "M", "N", "O"

# Columnas de Hitos
H_ID, H_PID, H_NOM, H_RESP, H_FEC, H_EST, H_ENT = "A", "B", "C", "D", "E", "F", "G"
H_PROY, H_DEP, H_SUB, H_DIAS, H_ALE = "H", "I", "J", "K", "L"


def load():
    tax = yaml.safe_load((DATA / "taxonomia.yaml").read_text(encoding="utf-8"))
    deps = yaml.safe_load((DATA / "departamentos.yaml").read_text(encoding="utf-8"))
    proys = yaml.safe_load((DATA / "proyectos.yaml").read_text(encoding="utf-8"))
    kpis = yaml.safe_load((DATA / "kpis.yaml").read_text(encoding="utf-8"))
    hitos = yaml.safe_load((DATA / "hitos.yaml").read_text(encoding="utf-8"))
    hitos.sort(key=lambda h: (h.get("fecha_objetivo") or dt.date(2100, 1, 1), h.get("id") or ""))
    return tax, deps, proys, kpis, hitos


def validate(tax, deps, proys, kpis, hitos):
    """Aborta antes de escribir si los datos son incoherentes."""
    errs = []
    pids = [p["id"] for p in proys]
    for dup in {i for i in pids if pids.count(i) > 1}:
        errs.append("id de proyecto duplicado: %s" % dup)
    subs = {s["nombre"] for d in deps for s in d["subareas"]}
    dnames = {d["nombre"] for d in deps}
    for p in proys:
        if p["departamento"] not in dnames:
            errs.append("%s: departamento desconocido '%s'" % (p["id"], p["departamento"]))
        if p["subarea"] not in subs:
            errs.append("%s: subárea desconocida '%s'" % (p["id"], p["subarea"]))
        if p["estado"] not in tax["estado_proyecto"]:
            errs.append("%s: estado desconocido '%s'" % (p["id"], p["estado"]))
        if p["fecha_fin_prevista"] < p["fecha_inicio"]:
            errs.append("%s: fin previsto anterior al inicio" % p["id"])
    for c, name in ((kpis, "KPI"), (hitos, "hito")):
        for x in c:
            if x["proyecto_id"] not in pids:
                errs.append("%s %s: proyecto_id inexistente '%s'" % (name, x["id"], x["proyecto_id"]))
    if errs:
        print("ERRORES EN LOS DATOS:", file=sys.stderr)
        for e in errs:
            print("  - " + e, file=sys.stderr)
        sys.exit(1)


# --------------------------------------------------------------------------
# Hoja Config — listas que alimentan los desplegables
# --------------------------------------------------------------------------
def build_config(wb, tax, deps, proys):
    ws = wb.create_sheet("Config")
    lists = [
        ("Estado proyecto", tax["estado_proyecto"]),
        ("Fase", tax["fase"]),
        ("Prioridad", tax["prioridad"]),
        ("Tipo", tax["tipo_proyecto"]),
        ("Riesgo", tax["riesgo"]),
        ("Estado hito", tax["estado_hito"]),
        ("Unidad KPI", tax["unidad_kpi"]),
        ("Dirección KPI", tax["direccion_kpi"]),
        ("Frecuencia KPI", tax["frecuencia_kpi"]),
        ("Fuente dato", tax["fuente_dato"]),
        ("Departamento", [d["nombre"] for d in deps]),
        ("Subárea", [s["nombre"] for d in deps for s in d["subareas"]]),
        ("Owner", sorted({p["owner"] for p in proys} |
                         {d["owner"] for d in deps} |
                         {s["owner"] for d in deps for s in d["subareas"]})),
        ("Estado activo", tax["estado_activo"]),
        ("Filtro departamento", ["(Todos)"] + [d["nombre"] for d in deps]),
        ("Filtro subárea", ["(Todas)"] + [s["nombre"] for d in deps for s in d["subareas"]]),
        ("Filtro estado", ["(Todos)"] + tax["estado_proyecto"]),
    ]
    for i, (name, values) in enumerate(lists):
        col = i + 1
        c = ws.cell(row=1, column=col, value=name)
        c.font = Font(name=FONT, size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=PURPLE)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        for j, v in enumerate(values):
            vc = ws.cell(row=2 + j, column=col, value=v)
            vc.font = Font(name=FONT, size=9, color=INK)
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = BORDER_C
    # Índice de columnas por nombre, para construir las validaciones
    idx = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(lists)}
    sizes = {name: len(values) for name, values in lists}
    return {n: "=Config!$%s$2:$%s$%d" % (idx[n], idx[n], 1 + sizes[n]) for n in idx}, idx


def add_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        errorStyle="warning",
                        error="Valor fuera de la lista. Si necesitas uno nuevo, añádelo en taxonomia.yaml y regenera.",
                        errorTitle="Valor no previsto")
    ws.add_data_validation(dv)
    dv.add(cells)


# --------------------------------------------------------------------------
# Hoja Proyectos
# --------------------------------------------------------------------------
PROY_HEADERS = [
    "ID", "Proyecto", "Departamento", "Subárea", "Owner", "Tipo", "Prioridad",
    "Estado", "Fase", "Inicio", "Fin previsto", "Avance %", "Riesgo",
    "Presupuesto €", "Gasto €", "Objetivo de negocio", "Impacto esperado", "Notas",
    "ƒ % ppto consumido", "ƒ Días restantes", "ƒ Avance esperado",
    "ƒ Desviación", "ƒ Semáforo", "ƒ KPIs", "ƒ % consecución KPIs",
    "ƒ Hitos", "ƒ Hitos OK", "ƒ Hitos vencidos",
]
PROY_NOTES = {
    "ƒ % ppto consumido": "Gasto € / Presupuesto €.",
    "ƒ Días restantes": "Fin previsto − hoy. Vacío si el proyecto está Completado o Cancelado.",
    "ƒ Avance esperado": "Proporción de calendario ya transcurrida, acotada entre 0 y 1: (hoy − inicio) / (fin − inicio).",
    "ƒ Desviación": "Avance real − avance esperado. Negativo = vamos por detrás del calendario.",
    "ƒ Semáforo": "Vencido si quedan días < 0; Inminente si quedan ≤ 14 días; Desviado si la desviación < −15pp; en el resto En plazo.",
    "ƒ % consecución KPIs": "Media del % de consecución de los KPIs del proyecto (hoja KPIs).",
    "ƒ Hitos vencidos": "Hitos con fecha objetivo pasada cuyo estado no es Completado ni Cancelado.",
}


def build_proyectos(wb, proys, dvs):
    ws = wb.create_sheet("Proyectos")
    table_header(ws, 1, 1, PROY_HEADERS, computed_from=19)
    for h, note in PROY_NOTES.items():
        col = PROY_HEADERS.index(h) + 1
        ws.cell(row=1, column=col).comment = Comment(note, "Growth OS")

    for i, p in enumerate(proys):
        r = P_FIRST + i
        ws.cell(row=r, column=1, value=p["id"])
        ws.cell(row=r, column=2, value=p["nombre"])
        ws.cell(row=r, column=3, value=p["departamento"])
        ws.cell(row=r, column=4, value=p["subarea"])
        ws.cell(row=r, column=5, value=p["owner"])
        ws.cell(row=r, column=6, value=p["tipo"])
        ws.cell(row=r, column=7, value=p["prioridad"])
        ws.cell(row=r, column=8, value=p["estado"])
        ws.cell(row=r, column=9, value=p["fase"])
        ws.cell(row=r, column=10, value=p["fecha_inicio"])
        ws.cell(row=r, column=11, value=p["fecha_fin_prevista"])
        ws.cell(row=r, column=12, value=p["avance_pct"])
        ws.cell(row=r, column=13, value=p["riesgo"])
        ws.cell(row=r, column=14, value=p.get("presupuesto_eur", 0))
        ws.cell(row=r, column=15, value=p.get("gasto_eur", 0))
        ws.cell(row=r, column=16, value=p.get("objetivo", ""))
        ws.cell(row=r, column=17, value=p.get("impacto_esperado", ""))
        ws.cell(row=r, column=18, value=p.get("notas", ""))

    activo = 'OR(${e}{r}="Completado",${e}{r}="Cancelado")'
    for r in range(P_FIRST, P_LAST + 1):
        g = 'IF($%s%d="","",' % (P_ID, r)
        cerrado = activo.format(e=P_EST, r=r)
        f = {
            P_CONS: '={g}IFERROR(${o}{r}/${n}{r},""))'.format(g=g, o=P_GASTO, n=P_PPTO, r=r),
            P_DIAS: '={g}IF({c},"",${k}{r}-TODAY()))'.format(g=g, c=cerrado, k=P_FIN, r=r),
            P_AVAE: '={g}IF({c},"",IFERROR(MEDIAN(0,(TODAY()-${j}{r})/(${k}{r}-${j}{r}),1),"")))'.format(
                g=g, c=cerrado, j=P_INI, k=P_FIN, r=r),
            P_DESV: '={g}IF(${u}{r}="","",${l}{r}-${u}{r}))'.format(g=g, u=P_AVAE, l=P_AVA, r=r),
            P_SEM: ('={g}IF({c},"—",IF(${t}{r}<0,"Vencido",IF(${t}{r}<=14,"Inminente",'
                    'IF(IF(${v}{r}="",FALSE,${v}{r}<-0.15),"Desviado","En plazo")))))').format(
                g=g, c=cerrado, t=P_DIAS, v=P_DESV, r=r),
            P_NKPI: '={g}COUNTIF({kb},$%s{r}))'.format(g=g, kb=K(K_PID), r=r) % P_ID,
            P_PKPI: '={g}IFERROR(AVERAGEIF({kb},$%s{r},{kp}),""))'.format(
                g=g, kb=K(K_PID), kp=K(K_PCT), r=r) % P_ID,
            P_NHIT: '={g}COUNTIF({hb},$%s{r}))'.format(g=g, hb=H(H_PID), r=r) % P_ID,
            P_HOK: '={g}COUNTIFS({hb},$%s{r},{he},"Completado"))'.format(
                g=g, hb=H(H_PID), he=H(H_EST), r=r) % P_ID,
            P_HVEN: ('={g}SUMPRODUCT(({hb}=$%s{r})*({hf}<>"")*({hf}<TODAY())'
                     '*({he}<>"Completado")*({he}<>"Cancelado")))').format(
                g=g, hb=H(H_PID), hf=H(H_FEC), he=H(H_EST), r=r) % P_ID,
        }
        for col, formula in f.items():
            ws["%s%d" % (col, r)] = formula

    # Formato
    for r in range(P_FIRST, P_LAST + 1):
        for col in range(1, len(PROY_HEADERS) + 1):
            c = ws.cell(row=r, column=col)
            L = get_column_letter(col)
            calc = col >= 19
            c.font = Font(name=FONT, size=9, color=GREY if calc else INK)
            c.border = thin()
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if L in ("B", "P", "Q", "R") else "center",
                                    wrap_text=L in ("P", "Q", "R"))
            if calc:
                c.fill = PatternFill("solid", fgColor=COMPUTED_BG)
            if L in (P_INI, P_FIN):
                c.number_format = FMT_DATE
            elif L in (P_PPTO, P_GASTO):
                c.number_format = FMT_EUR
            elif L == P_AVA:
                c.number_format = FMT_PCT_IN
            elif L in (P_CONS, P_AVAE, P_PKPI):
                c.number_format = FMT_PCT0
            elif L == P_DESV:
                c.number_format = '+0%;-0%;0%'
            elif L in (P_DIAS, P_NKPI, P_NHIT, P_HOK, P_HVEN):
                c.number_format = FMT_INT
        ws.row_dimensions[r].height = 17

    widths(ws, {"A": 10, "B": 40, "C": 13, "D": 22, "E": 16, "F": 15, "G": 9,
                "H": 12, "I": 11, "J": 11, "K": 11, "L": 10, "M": 9, "N": 14,
                "O": 13, "P": 42, "Q": 42, "R": 42, "S": 12, "T": 11, "U": 12,
                "V": 11, "W": 12, "X": 8, "Y": 13, "Z": 8, "AA": 9, "AB": 12})
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(PROY_HEADERS)), P_LAST)
    ws.sheet_properties.tabColor = PURPLE

    # Validaciones
    add_dv(ws, dvs["Departamento"], "%s%d:%s%d" % (P_DEP, P_FIRST, P_DEP, P_LAST))
    add_dv(ws, dvs["Subárea"], "%s%d:%s%d" % (P_SUB, P_FIRST, P_SUB, P_LAST))
    add_dv(ws, dvs["Owner"], "%s%d:%s%d" % (P_OWN, P_FIRST, P_OWN, P_LAST))
    add_dv(ws, dvs["Tipo"], "%s%d:%s%d" % (P_TIPO, P_FIRST, P_TIPO, P_LAST))
    add_dv(ws, dvs["Prioridad"], "%s%d:%s%d" % (P_PRIO, P_FIRST, P_PRIO, P_LAST))
    add_dv(ws, dvs["Estado proyecto"], "%s%d:%s%d" % (P_EST, P_FIRST, P_EST, P_LAST))
    add_dv(ws, dvs["Fase"], "%s%d:%s%d" % (P_FASE, P_FIRST, P_FASE, P_LAST))
    add_dv(ws, dvs["Riesgo"], "%s%d:%s%d" % (P_RIE, P_FIRST, P_RIE, P_LAST))

    # Formato condicional
    state_rules(ws, "%s%d:%s%d" % (P_EST, P_FIRST, P_EST, P_LAST), ESTADO_COLORS)
    state_rules(ws, "%s%d:%s%d" % (P_SEM, P_FIRST, P_SEM, P_LAST), SEMAFORO_COLORS)
    state_rules(ws, "%s%d:%s%d" % (P_RIE, P_FIRST, P_RIE, P_LAST), RIESGO_COLORS)
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (P_AVA, P_FIRST, P_AVA, P_LAST),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=PURPLE))
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (P_PKPI, P_FIRST, P_PKPI, P_LAST),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A3D9C5"))
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (P_DESV, P_FIRST, P_DESV, P_LAST),
        CellIsRule(operator="lessThan", formula=["-0.15"],
                   font=Font(name=FONT, size=9, bold=True, color=RED_FG),
                   fill=PatternFill("solid", bgColor=RED_BG)))
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (P_HVEN, P_FIRST, P_HVEN, P_LAST),
        CellIsRule(operator="greaterThan", formula=["0"],
                   font=Font(name=FONT, size=9, bold=True, color=RED_FG),
                   fill=PatternFill("solid", bgColor=RED_BG)))
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (P_CONS, P_FIRST, P_CONS, P_LAST),
        CellIsRule(operator="greaterThan", formula=["1"],
                   font=Font(name=FONT, size=9, bold=True, color=RED_FG),
                   fill=PatternFill("solid", bgColor=RED_BG)))
    return ws


# --------------------------------------------------------------------------
# Hoja KPIs
# --------------------------------------------------------------------------
KPI_HEADERS = [
    "ID", "Proyecto ID", "KPI", "Unidad", "Dirección", "Baseline", "Target",
    "Actual", "Frecuencia", "Fuente",
    "ƒ Proyecto", "ƒ Departamento", "ƒ Subárea", "ƒ % consecución", "ƒ Estado",
]
KPI_NOTES = {
    "Dirección": "Subir = mejor cuanto más alto (matrículas, CVR). Bajar = mejor cuanto más bajo (CAC, CPL, días de ciclo).",
    "ƒ % consecución": ("Si Dirección = Subir: (Actual − Baseline) / (Target − Baseline). "
                        "Si Dirección = Bajar: (Baseline − Actual) / (Baseline − Target). "
                        "100% = target alcanzado; puede pasar de 100% o ser negativo."),
    "ƒ Estado": "≥100% Conseguido · ≥70% En camino · ≥30% Atención · <30% Crítico.",
}
UNIT_FMT = {"EUR": FMT_EUR, "%": FMT_PCT1, "nº": FMT_INT, "ratio": '0.00"x"', "días": '0.# "d"'}


def build_kpis(wb, kpis, dvs):
    ws = wb.create_sheet("KPIs")
    table_header(ws, 1, 1, KPI_HEADERS, computed_from=11)
    for h, note in KPI_NOTES.items():
        ws.cell(row=1, column=KPI_HEADERS.index(h) + 1).comment = Comment(note, "Growth OS")

    for i, k in enumerate(kpis):
        r = K_FIRST + i
        ws.cell(row=r, column=1, value=k["id"])
        ws.cell(row=r, column=2, value=k["proyecto_id"])
        ws.cell(row=r, column=3, value=k["nombre"])
        ws.cell(row=r, column=4, value=k["unidad"])
        ws.cell(row=r, column=5, value=k["direccion"])
        ws.cell(row=r, column=6, value=k["baseline"])
        ws.cell(row=r, column=7, value=k["target"])
        ws.cell(row=r, column=8, value=k["actual"])
        ws.cell(row=r, column=9, value=k["frecuencia"])
        ws.cell(row=r, column=10, value=k["fuente"])
        fmt = UNIT_FMT.get(k["unidad"], FMT_NUM)
        for col in (6, 7, 8):
            ws.cell(row=r, column=col).number_format = fmt

    nf = 'IF($%s{r}="","",' % K_PID
    for r in range(K_FIRST, K_LAST + 1):
        lookup = 'IFERROR(INDEX({src},MATCH($%s{r},{ids},0)),"⚠ ID no encontrado")' % K_PID
        ws["%s%d" % (K_PROY, r)] = "=" + nf.format(r=r) + lookup.format(
            src=P(P_NOM), ids=P(P_ID), r=r) + ")"
        ws["%s%d" % (K_DEP, r)] = "=" + nf.format(r=r) + lookup.format(
            src=P(P_DEP), ids=P(P_ID), r=r) + ")"
        ws["%s%d" % (K_SUB, r)] = "=" + nf.format(r=r) + lookup.format(
            src=P(P_SUB), ids=P(P_ID), r=r) + ")"
        ws["%s%d" % (K_PCT, r)] = (
            '=IF(OR($%s{r}="",$%s{r}="",$%s{r}="",$%s{r}="",$%s{r}=""),"",'
            'IFERROR(IF($%s{r}="Bajar",($%s{r}-$%s{r})/($%s{r}-$%s{r}),'
            '($%s{r}-$%s{r})/($%s{r}-$%s{r})),""))' % (
                K_PID, K_DIR, K_BASE, K_TGT, K_ACT,
                K_DIR, K_BASE, K_ACT, K_BASE, K_TGT,
                K_ACT, K_BASE, K_TGT, K_BASE)).format(r=r)
        ws["%s%d" % (K_EST, r)] = (
            '=IF($%s{r}="","",IF($%s{r}>=1,"Conseguido",IF($%s{r}>=0.7,"En camino",'
            'IF($%s{r}>=0.3,"Atención","Crítico"))))' % (K_PCT, K_PCT, K_PCT, K_PCT)).format(r=r)

    for r in range(K_FIRST, K_LAST + 1):
        for col in range(1, len(KPI_HEADERS) + 1):
            c = ws.cell(row=r, column=col)
            L = get_column_letter(col)
            calc = col >= 11
            c.font = Font(name=FONT, size=9, color=GREY if calc else INK)
            c.border = thin()
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if L in ("C", "K") else "center")
            if calc:
                c.fill = PatternFill("solid", fgColor=COMPUTED_BG)
            if L == K_PCT:
                c.number_format = FMT_PCT0
            elif L in (K_BASE, K_TGT, K_ACT) and not c.number_format.startswith("#"):
                if c.value is None:
                    c.number_format = FMT_NUM
        ws.row_dimensions[r].height = 17

    widths(ws, {"A": 10, "B": 12, "C": 34, "D": 10, "E": 11, "F": 12, "G": 12,
                "H": 12, "I": 12, "J": 15, "K": 34, "L": 13, "M": 22, "N": 14, "O": 12})
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(KPI_HEADERS)), K_LAST)
    ws.sheet_properties.tabColor = PURPLE

    add_dv(ws, "=%s" % P(P_ID), "%s%d:%s%d" % (K_PID, K_FIRST, K_PID, K_LAST))
    add_dv(ws, dvs["Unidad KPI"], "%s%d:%s%d" % (K_UNI, K_FIRST, K_UNI, K_LAST))
    add_dv(ws, dvs["Dirección KPI"], "%s%d:%s%d" % (K_DIR, K_FIRST, K_DIR, K_LAST))
    add_dv(ws, dvs["Frecuencia KPI"], "%s%d:%s%d" % (K_FREQ, K_FIRST, K_FREQ, K_LAST))
    add_dv(ws, dvs["Fuente dato"], "%s%d:%s%d" % (K_SRC, K_FIRST, K_SRC, K_LAST))

    state_rules(ws, "%s%d:%s%d" % (K_EST, K_FIRST, K_EST, K_LAST), KPI_COLORS)
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (K_PCT, K_FIRST, K_PCT, K_LAST),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=PURPLE))
    return ws


# --------------------------------------------------------------------------
# Hoja Hitos
# --------------------------------------------------------------------------
HITO_HEADERS = [
    "ID", "Proyecto ID", "Hito", "Responsable", "Fecha objetivo", "Estado", "Entregable",
    "ƒ Proyecto", "ƒ Departamento", "ƒ Subárea", "ƒ Días restantes", "ƒ Alerta",
]
HITO_NOTES = {
    "ƒ Días restantes": "Fecha objetivo − hoy. Vacío si el hito está Completado o Cancelado.",
    "ƒ Alerta": "Vencido (<0 días) · Esta semana (≤7) · Próximos 30 días (≤30) · En plazo.",
}


def build_hitos(wb, hitos, dvs):
    ws = wb.create_sheet("Hitos")
    table_header(ws, 1, 1, HITO_HEADERS, computed_from=8)
    for h, note in HITO_NOTES.items():
        ws.cell(row=1, column=HITO_HEADERS.index(h) + 1).comment = Comment(note, "Growth OS")

    for i, h in enumerate(hitos):
        r = H_FIRST + i
        ws.cell(row=r, column=1, value=h["id"])
        ws.cell(row=r, column=2, value=h["proyecto_id"])
        ws.cell(row=r, column=3, value=h["nombre"])
        ws.cell(row=r, column=4, value=h["responsable"])
        ws.cell(row=r, column=5, value=h["fecha_objetivo"])
        ws.cell(row=r, column=6, value=h["estado"])
        ws.cell(row=r, column=7, value=h.get("entregable", ""))

    nf = 'IF($%s{r}="","",' % H_PID
    for r in range(H_FIRST, H_LAST + 1):
        lookup = 'IFERROR(INDEX({src},MATCH($%s{r},{ids},0)),"⚠ ID no encontrado")' % H_PID
        ws["%s%d" % (H_PROY, r)] = "=" + nf.format(r=r) + lookup.format(
            src=P(P_NOM), ids=P(P_ID), r=r) + ")"
        ws["%s%d" % (H_DEP, r)] = "=" + nf.format(r=r) + lookup.format(
            src=P(P_DEP), ids=P(P_ID), r=r) + ")"
        ws["%s%d" % (H_SUB, r)] = "=" + nf.format(r=r) + lookup.format(
            src=P(P_SUB), ids=P(P_ID), r=r) + ")"
        cerrado = 'OR($%s{r}="Completado",$%s{r}="Cancelado")'.format(r=r) % (H_EST, H_EST)
        cerrado = cerrado.format(r=r)
        ws["%s%d" % (H_DIAS, r)] = (
            '=IF($%s%d="","",IF(%s,"",IF($%s%d="","",$%s%d-TODAY())))' % (
                H_ID, r, cerrado, H_FEC, r, H_FEC, r))
        ws["%s%d" % (H_ALE, r)] = (
            '=IF($%s%d="","",IF(%s,"—",IF($%s%d="","",IF($%s%d<0,"Vencido",'
            'IF($%s%d<=7,"Esta semana",IF($%s%d<=30,"Próximos 30 días","En plazo"))))))' % (
                H_ID, r, cerrado, H_DIAS, r, H_DIAS, r, H_DIAS, r, H_DIAS, r))

    for r in range(H_FIRST, H_LAST + 1):
        for col in range(1, len(HITO_HEADERS) + 1):
            c = ws.cell(row=r, column=col)
            L = get_column_letter(col)
            calc = col >= 8
            c.font = Font(name=FONT, size=9, color=GREY if calc else INK)
            c.border = thin()
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if L in ("C", "G", "H") else "center",
                                    wrap_text=(L == "G"))
            if calc:
                c.fill = PatternFill("solid", fgColor=COMPUTED_BG)
            if L == H_FEC:
                c.number_format = FMT_DATE
            elif L == H_DIAS:
                c.number_format = FMT_INT
        ws.row_dimensions[r].height = 17

    widths(ws, {"A": 9, "B": 12, "C": 44, "D": 16, "E": 14, "F": 13, "G": 40,
                "H": 34, "I": 13, "J": 22, "K": 13, "L": 17})
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(HITO_HEADERS)), H_LAST)
    ws.sheet_properties.tabColor = PURPLE

    add_dv(ws, "=%s" % P(P_ID), "%s%d:%s%d" % (H_PID, H_FIRST, H_PID, H_LAST))
    add_dv(ws, dvs["Owner"], "%s%d:%s%d" % (H_RESP, H_FIRST, H_RESP, H_LAST))
    add_dv(ws, dvs["Estado hito"], "%s%d:%s%d" % (H_EST, H_FIRST, H_EST, H_LAST))

    state_rules(ws, "%s%d:%s%d" % (H_EST, H_FIRST, H_EST, H_LAST), {
        "Pendiente": (NEUTRAL_BG, NEUTRAL_FG), "En curso": (PURPLE_BG, PURPLE_FG),
        "Completado": (GREEN_BG, GREEN_FG), "Retrasado": (RED_BG, RED_FG),
        "Cancelado": (NEUTRAL_BG, NEUTRAL_FG)})
    state_rules(ws, "%s%d:%s%d" % (H_ALE, H_FIRST, H_ALE, H_LAST), ALERTA_COLORS)
    return ws


# --------------------------------------------------------------------------
# Listas filtradas dinámicas (equivalente a FILTER sin funciones de matriz:
# dos columnas auxiliares de conteo + INDEX/MATCH)
# --------------------------------------------------------------------------
def dynamic_list(ws, cond_tpl, src_first, src_last, id_col, cnt_col, out_col,
                 disp_first, disp_rows):
    for r in range(src_first, src_last + 1):
        cond = cond_tpl.format(r=r)
        prev = "%s%d" % (cnt_col, r - 1)
        ws["%s%d" % (cnt_col, r)] = "=IF(%s,N(%s)+1,N(%s))" % (cond, prev, prev)
        ws["%s%d" % (id_col, r)] = '=IF(%s,%s%d,"")' % (cond, cnt_col, r)
    for i in range(disp_rows):
        ws["%s%d" % (out_col, disp_first + i)] = '=IFERROR(MATCH(%d,$%s$%d:$%s$%d,0),"")' % (
            i + 1, id_col, src_first, id_col, src_last)


def dyn(out_col, row, src_range):
    return '=IF($%s%d="","",INDEX(%s,$%s%d))' % (out_col, row, src_range, out_col, row)


def fill_dyn_table(ws, out_col, first_row, n_rows, cols, first_col=2):
    """cols = lista de (rango_origen, number_format o None, align)."""
    for i in range(n_rows):
        r = first_row + i
        for j, (src, fmt, align) in enumerate(cols):
            c = ws.cell(row=r, column=first_col + j, value=dyn(out_col, r, src))
            c.font = Font(name=FONT, size=9, color=INK)
            c.border = thin()
            c.alignment = Alignment(vertical="center", horizontal=align)
            if fmt:
                c.number_format = fmt
        ws.row_dimensions[r].height = 17


def hide_helpers(ws, letters):
    for L in letters:
        ws.column_dimensions[L].hidden = True
        ws.column_dimensions[L].width = 3


# --------------------------------------------------------------------------
# Hoja Dashboard (nivel compañía)
# --------------------------------------------------------------------------
def build_dashboard(wb, tax, deps, cfg_idx):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PURPLE_DARK

    title(ws, "B2", "Founderz · Growth OS")
    subtitle(ws, "B3", "Panel de dirección — portfolio de proyectos, KPIs e hitos. "
                       "Todo lo que ves aquí se calcula desde las hojas Proyectos, KPIs y Hitos.")
    ws["B4"] = "=\"Datos a fecha de \"&TEXT(TODAY(),\"dd/mm/yyyy\")"
    ws["B4"].font = Font(name=FONT, size=9, italic=True, color=PURPLE)

    no_vacio = '({a}<>"")'.format(a=P(P_ID))
    no_canc = '({e}<>"Cancelado")'.format(e=P(P_EST))
    activos = "+".join('COUNTIF(%s,"%s")' % (P(P_EST), e) for e in tax["estado_activo"])
    hitos_venc = ('SUMPRODUCT(({a}<>"")*({f}<>"")*({f}<TODAY())*({e}<>"Completado")'
                  '*({e}<>"Cancelado"))').format(a=H(H_ID), f=H(H_FEC), e=H(H_EST))

    tiles = [
        ("Proyectos totales", "=COUNTA(%s)" % P(P_ID), FMT_INT,
         "Filas con ID en la hoja Proyectos."),
        ("Activos", "=" + activos, FMT_INT,
         "Estados considerados activos: " + ", ".join(tax["estado_activo"]) + "."),
        ("En riesgo o bloqueados", '=COUNTIF(%s,"En riesgo")+COUNTIF(%s,"Bloqueado")' % (
            P(P_EST), P(P_EST)), FMT_INT, None),
        ("Avance medio", '=IFERROR(SUMPRODUCT(%s*%s*%s)/SUMPRODUCT(%s*%s*1),"")' % (
            no_vacio, no_canc, P(P_AVA), no_vacio, no_canc), FMT_PCT0,
         "Media simple del avance de los proyectos no cancelados."),
        ("Presupuesto comprometido", "=SUMPRODUCT(%s*%s)" % (no_canc, P(P_PPTO)), FMT_EUR,
         "Suma del presupuesto de los proyectos no cancelados."),
        ("Gasto acumulado", "=SUMPRODUCT(%s*%s)" % (no_canc, P(P_GASTO)), FMT_EUR, None),
        ("% presupuesto consumido", '=IFERROR(SUMPRODUCT(%s*%s)/SUMPRODUCT(%s*%s),"")' % (
            no_canc, P(P_GASTO), no_canc, P(P_PPTO)), FMT_PCT0, None),
        ("Hitos vencidos", "=" + hitos_venc, FMT_INT,
         "Hitos con fecha objetivo pasada y estado distinto de Completado o Cancelado."),
        ("KPIs en crítico", '=COUNTIF(%s,"Crítico")' % K(K_EST), FMT_INT,
         "KPIs con menos del 30% del recorrido baseline→target cubierto."),
    ]
    for i, (label, formula, fmt, note) in enumerate(tiles):
        tile(ws, 6, 2 + i, label, formula, fmt, note)

    row = 10

    # --- Portfolio por estado ---
    section(ws, row, 2, "Portfolio por estado")
    row += 1
    hdr = ["Estado", "Proyectos", "% del total", "Presupuesto €", "Gasto €", "Avance medio"]
    table_header(ws, row, 2, hdr)
    first = row + 1
    last = first + len(tax["estado_proyecto"]) - 1
    for i, e in enumerate(tax["estado_proyecto"]):
        r = first + i
        vals = [
            e,
            "=COUNTIF(%s,$B%d)" % (P(P_EST), r),
            '=IFERROR($C%d/SUM($C$%d:$C$%d),"")' % (r, first, last),
            "=SUMIF(%s,$B%d,%s)" % (P(P_EST), r, P(P_PPTO)),
            "=SUMIF(%s,$B%d,%s)" % (P(P_EST), r, P(P_GASTO)),
            '=IFERROR(AVERAGEIF(%s,$B%d,%s),"")' % (P(P_EST), r, P(P_AVA)),
        ]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + j, value=v)
            c.font = Font(name=FONT, size=9, color=INK)
            c.border = thin()
            c.alignment = Alignment(vertical="center", horizontal="left" if j == 0 else "center")
            nf_e = [None, FMT_INT, FMT_PCT0, FMT_EUR, FMT_EUR, FMT_PCT0][j]
            if nf_e:
                c.number_format = nf_e
    state_rules(ws, "B%d:B%d" % (first, last), ESTADO_COLORS)
    row = last + 2

    # --- Portfolio por departamento y por subárea ---
    def agg_block(caption, entries, with_dep_col, sub_level):
        nonlocal row
        section(ws, row, 2, caption)
        row += 1
        base = (["Departamento", "Subárea"] if with_dep_col else ["Departamento"]) + [
            "Proyectos", "Activos", "Presupuesto €", "Gasto €", "% consumido",
            "Avance medio", "Hitos vencidos", "KPIs críticos"]
        table_header(ws, row, 2, base)
        f = row + 1
        keycol = get_column_letter(3 if with_dep_col else 2)
        pcol = P(P_SUB) if sub_level else P(P_DEP)
        kcol = K(K_SUB) if sub_level else K(K_DEP)
        hcol = H(H_SUB) if sub_level else H(H_DEP)
        for i, ent in enumerate(entries):
            r = f + i
            key = "$%s%d" % (keycol, r)
            act = "+".join('COUNTIFS(%s,%s,%s,"%s")' % (pcol, key, P(P_EST), e)
                           for e in tax["estado_activo"])
            ncols = 2 if with_dep_col else 1
            vals = (list(ent) if with_dep_col else [ent[0]]) + [
                "=COUNTIF(%s,%s)" % (pcol, key),
                "=" + act,
                "=SUMIF(%s,%s,%s)" % (pcol, key, P(P_PPTO)),
                "=SUMIF(%s,%s,%s)" % (pcol, key, P(P_GASTO)),
                '=IFERROR(%s%d/%s%d,"")' % (get_column_letter(2 + ncols + 3), r,
                                            get_column_letter(2 + ncols + 2), r),
                '=IFERROR(AVERAGEIF(%s,%s,%s),"")' % (pcol, key, P(P_AVA)),
                ('=SUMPRODUCT((%s=%s)*(%s<>"")*(%s<TODAY())*(%s<>"Completado")'
                 '*(%s<>"Cancelado"))') % (hcol, key, H(H_FEC), H(H_FEC), H(H_EST), H(H_EST)),
                '=COUNTIFS(%s,%s,%s,"Crítico")' % (kcol, key, K(K_EST)),
            ]
            fmts = [None] * ncols + [FMT_INT, FMT_INT, FMT_EUR, FMT_EUR, FMT_PCT0,
                                     FMT_PCT0, FMT_INT, FMT_INT]
            for j, v in enumerate(vals):
                c = ws.cell(row=r, column=2 + j, value=v)
                c.font = Font(name=FONT, size=9, color=INK)
                c.border = thin()
                c.alignment = Alignment(vertical="center",
                                        horizontal="left" if j < ncols else "center")
                if fmts[j]:
                    c.number_format = fmts[j]
            for cc in (2 + ncols + 6, 2 + ncols + 7):
                ws.conditional_formatting.add(
                    "%s%d" % (get_column_letter(cc), r),
                    CellIsRule(operator="greaterThan", formula=["0"],
                               font=Font(name=FONT, size=9, bold=True, color=RED_FG),
                               fill=PatternFill("solid", bgColor=RED_BG)))
        row = f + len(entries) + 1

    agg_block("Portfolio por departamento", [(d["nombre"],) for d in deps], False, False)
    agg_block("Zoom por subárea", [(d["nombre"], s["nombre"]) for d in deps
                                   for s in d["subareas"]], True, True)

    # --- Proyectos que requieren atención ---
    section(ws, row, 2, "Proyectos que requieren atención")
    subtitle(ws, "F%d" % row, "En riesgo · bloqueados · con hitos vencidos · desviados >15pp", 8)
    row += 1
    table_header(ws, row, 2, ["ID", "Proyecto", "Subárea", "Owner", "Prioridad", "Estado",
                              "Fin previsto", "Avance", "Desviación", "Semáforo",
                              "Hitos vencidos", "% KPIs"])
    att_first = row + 1
    att_rows = 20
    cond_att = ('AND({a}<>"",OR({e}="En riesgo",{e}="Bloqueado",{hv}>0,'
                'IF({dv}="",FALSE,{dv}<-0.15)))').format(
        a="Proyectos!$%s{r}" % P_ID, e="Proyectos!$%s{r}" % P_EST,
        hv="Proyectos!$%s{r}" % P_HVEN, dv="Proyectos!$%s{r}" % P_DESV)
    dynamic_list(ws, cond_att, P_FIRST, P_LAST, "BJ", "BK", "BH", att_first, att_rows)
    fill_dyn_table(ws, "BH", att_first, att_rows, [
        (P(P_ID), None, "center"), (P(P_NOM), None, "left"), (P(P_SUB), None, "left"),
        (P(P_OWN), None, "left"), (P(P_PRIO), None, "center"), (P(P_EST), None, "center"),
        (P(P_FIN), FMT_DATE, "center"), (P(P_AVA), FMT_PCT0, "center"),
        (P(P_DESV), '+0%;-0%;0%', "center"), (P(P_SEM), None, "center"),
        (P(P_HVEN), FMT_INT, "center"), (P(P_PKPI), FMT_PCT0, "center")])
    state_rules(ws, "G%d:G%d" % (att_first, att_first + att_rows - 1), ESTADO_COLORS)
    state_rules(ws, "K%d:K%d" % (att_first, att_first + att_rows - 1), SEMAFORO_COLORS)
    row = att_first + att_rows + 1

    # --- Hitos vencidos y próximos 30 días ---
    section(ws, row, 2, "Hitos vencidos y de los próximos 30 días")
    row += 1
    table_header(ws, row, 2, ["Fecha objetivo", "Alerta", "Hito", "Proyecto",
                              "Subárea", "Responsable", "Estado"])
    hit_first = row + 1
    hit_rows = 25
    cond_h = ('AND({a}<>"",OR({al}="Vencido",{al}="Esta semana",{al}="Próximos 30 días"))'
              ).format(a="Hitos!$%s{r}" % H_ID, al="Hitos!$%s{r}" % H_ALE)
    dynamic_list(ws, cond_h, H_FIRST, H_LAST, "BN", "BO", "BL", hit_first, hit_rows)
    fill_dyn_table(ws, "BL", hit_first, hit_rows, [
        (H(H_FEC), FMT_DATE, "center"), (H(H_ALE), None, "center"),
        (H(H_NOM), None, "left"), (H(H_PROY), None, "left"), (H(H_SUB), None, "left"),
        (H(H_RESP), None, "left"), (H(H_EST), None, "center")])
    state_rules(ws, "C%d:C%d" % (hit_first, hit_first + hit_rows - 1), ALERTA_COLORS)

    widths(ws, {"A": 3, "B": 17, "C": 34, "D": 22, "E": 18, "F": 15, "G": 15,
                "H": 14, "I": 13, "J": 13, "K": 13, "L": 13, "M": 13})
    hide_helpers(ws, ["BH", "BJ", "BK", "BL", "BN", "BO"])
    ws.freeze_panes = "B6"
    return ws


# --------------------------------------------------------------------------
# Hoja Zoom departamento (nivel departamento / subárea)
# --------------------------------------------------------------------------
FDEP, FSUB, FEST = "$C$5", "$C$6", "$C$7"


def any_of(cell, all_value, range_):
    """Array 0/1: verdadero si el filtro está en '(Todos)' o si coincide."""
    return '(({c}="{a}")+({r}={c})>0)'.format(c=cell, a=all_value, r=range_)


def build_zoom(wb, tax, dvs):
    ws = wb.create_sheet("Zoom departamento")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PURPLE_DARK

    title(ws, "B2", "Zoom por departamento")
    subtitle(ws, "B3", "Elige departamento, subárea y estado en los desplegables: las tres tablas "
                       "de abajo y las tarjetas se recalculan solas.")

    for r, (label, dv_name, cell) in enumerate([
            ("Departamento", "Filtro departamento", "C5"),
            ("Subárea", "Filtro subárea", "C6"),
            ("Estado", "Filtro estado", "C7")], start=5):
        lc = ws.cell(row=r, column=2, value=label)
        lc.font = Font(name=FONT, size=9, bold=True, color=PURPLE)
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc = ws[cell]
        vc.value = "(Todos)" if label != "Subárea" else "(Todas)"
        vc.font = Font(name=FONT, size=11, bold=True, color=INK)
        vc.fill = PatternFill("solid", fgColor="FFF9DB")
        vc.border = thin(PURPLE)
        vc.alignment = Alignment(horizontal="left", vertical="center")
        add_dv(ws, dvs[dv_name], "%s:%s" % (cell, cell))
        ws.row_dimensions[r].height = 20
    ws.merge_cells("C5:E5")
    ws.merge_cells("C6:E6")
    ws.merge_cells("C7:E7")
    subtitle(ws, "G5", "Celdas amarillas = filtros editables. El resto de la hoja es fórmula.", 8)

    mdep = any_of(FDEP, "(Todos)", P(P_DEP))
    msub = any_of(FSUB, "(Todas)", P(P_SUB))
    mest = any_of(FEST, "(Todos)", P(P_EST))
    base = '({a}<>"")*{d}*{s}*{e}'.format(a=P(P_ID), d=mdep, s=msub, e=mest)
    act = "+".join('({e}="{v}")'.format(e=P(P_EST), v=v) for v in tax["estado_activo"])
    base_h = '({a}<>"")*{d}*{s}'.format(a=H(H_ID),
                                        d=any_of(FDEP, "(Todos)", H(H_DEP)),
                                        s=any_of(FSUB, "(Todas)", H(H_SUB)))
    base_k = '({a}<>"")*{d}*{s}'.format(a=K(K_ID),
                                        d=any_of(FDEP, "(Todos)", K(K_DEP)),
                                        s=any_of(FSUB, "(Todas)", K(K_SUB)))

    tiles = [
        ("Proyectos en el filtro", "=SUMPRODUCT(%s)" % base, FMT_INT, None),
        ("Activos", "=SUMPRODUCT(%s*((%s)>0))" % (base, act), FMT_INT,
         "Estados activos: " + ", ".join(tax["estado_activo"]) + "."),
        ("En riesgo o bloqueados", '=SUMPRODUCT(%s*((({e}="En riesgo")+({e}="Bloqueado"))>0))'.format(
            e=P(P_EST)) % base, FMT_INT, None),
        ("Avance medio", '=IFERROR(SUMPRODUCT(%s*%s)/SUMPRODUCT(%s),"")' % (
            base, P(P_AVA), base), FMT_PCT0, None),
        ("Presupuesto", "=SUMPRODUCT(%s*%s)" % (base, P(P_PPTO)), FMT_EUR, None),
        ("Gasto", "=SUMPRODUCT(%s*%s)" % (base, P(P_GASTO)), FMT_EUR, None),
        ("% consumido", '=IFERROR(SUMPRODUCT(%s*%s)/SUMPRODUCT(%s*%s),"")' % (
            base, P(P_GASTO), base, P(P_PPTO)), FMT_PCT0, None),
        ("Hitos vencidos", '=SUMPRODUCT(%s*({f}<>"")*({f}<TODAY())*({e}<>"Completado")*({e}<>"Cancelado"))'.format(
            f=H(H_FEC), e=H(H_EST)) % base_h, FMT_INT,
         "Los hitos y los KPIs se filtran por departamento y subárea; el filtro de Estado "
         "sólo se aplica a la tabla de proyectos."),
        ("KPIs en crítico", '=SUMPRODUCT(%s*({e}="Crítico"))'.format(e=K(K_EST)) % base_k,
         FMT_INT, None),
    ]
    for i, (label, formula, fmt, note) in enumerate(tiles):
        tile(ws, 9, 2 + i, label, formula, fmt, note)

    # --- Proyectos del filtro ---
    section(ws, 12, 2, "Proyectos del filtro")
    table_header(ws, 13, 2, ["ID", "Proyecto", "Subárea", "Owner", "Prioridad", "Estado",
                             "Fase", "Fin previsto", "Avance", "Esperado", "Desviación",
                             "Semáforo", "Presupuesto €", "Gasto €", "% ppto",
                             "% KPIs", "Hitos vencidos"])
    p_first, p_rows = 14, 40
    cond_p = ('AND({a}<>"",OR({fd}="(Todos)",{d}={fd}),OR({fs}="(Todas)",{s}={fs}),'
              'OR({fe}="(Todos)",{e}={fe}))').format(
        a="Proyectos!$%s{r}" % P_ID, d="Proyectos!$%s{r}" % P_DEP,
        s="Proyectos!$%s{r}" % P_SUB, e="Proyectos!$%s{r}" % P_EST,
        fd=FDEP, fs=FSUB, fe=FEST)
    dynamic_list(ws, cond_p, P_FIRST, P_LAST, "BJ", "BK", "BH", p_first, p_rows)
    fill_dyn_table(ws, "BH", p_first, p_rows, [
        (P(P_ID), None, "center"), (P(P_NOM), None, "left"), (P(P_SUB), None, "left"),
        (P(P_OWN), None, "left"), (P(P_PRIO), None, "center"), (P(P_EST), None, "center"),
        (P(P_FASE), None, "center"), (P(P_FIN), FMT_DATE, "center"),
        (P(P_AVA), FMT_PCT0, "center"), (P(P_AVAE), FMT_PCT0, "center"),
        (P(P_DESV), '+0%;-0%;0%', "center"), (P(P_SEM), None, "center"),
        (P(P_PPTO), FMT_EUR, "right"), (P(P_GASTO), FMT_EUR, "right"),
        (P(P_CONS), FMT_PCT0, "center"), (P(P_PKPI), FMT_PCT0, "center"),
        (P(P_HVEN), FMT_INT, "center")])
    state_rules(ws, "G%d:G%d" % (p_first, p_first + p_rows - 1), ESTADO_COLORS)
    state_rules(ws, "M%d:M%d" % (p_first, p_first + p_rows - 1), SEMAFORO_COLORS)
    ws.conditional_formatting.add(
        "J%d:J%d" % (p_first, p_first + p_rows - 1),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=PURPLE))

    # --- KPIs del filtro ---
    k_sec = p_first + p_rows + 1
    section(ws, k_sec, 2, "KPIs de los proyectos del filtro")
    table_header(ws, k_sec + 1, 2, ["Proyecto", "KPI", "Unidad", "Dirección", "Baseline",
                                    "Target", "Actual", "% consecución", "Estado",
                                    "Frecuencia", "Fuente"])
    k_first, k_rows = k_sec + 2, 40
    cond_k = ('AND({a}<>"",OR({fd}="(Todos)",{d}={fd}),OR({fs}="(Todas)",{s}={fs}))').format(
        a="KPIs!$%s{r}" % K_ID, d="KPIs!$%s{r}" % K_DEP, s="KPIs!$%s{r}" % K_SUB,
        fd=FDEP, fs=FSUB)
    dynamic_list(ws, cond_k, K_FIRST, K_LAST, "BN", "BO", "BL", k_first, k_rows)
    fill_dyn_table(ws, "BL", k_first, k_rows, [
        (K(K_PROY), None, "left"), (K(K_NOM), None, "left"), (K(K_UNI), None, "center"),
        (K(K_DIR), None, "center"), (K(K_BASE), FMT_NUM, "right"),
        (K(K_TGT), FMT_NUM, "right"), (K(K_ACT), FMT_NUM, "right"),
        (K(K_PCT), FMT_PCT0, "center"), (K(K_EST), None, "center"),
        (K(K_FREQ), None, "center"), (K(K_SRC), None, "center")])
    state_rules(ws, "J%d:J%d" % (k_first, k_first + k_rows - 1), KPI_COLORS)
    ws.conditional_formatting.add(
        "I%d:I%d" % (k_first, k_first + k_rows - 1),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=PURPLE))

    # --- Hitos del filtro ---
    h_sec = k_first + k_rows + 1
    section(ws, h_sec, 2, "Hitos vencidos y de los próximos 30 días")
    table_header(ws, h_sec + 1, 2, ["Fecha objetivo", "Alerta", "Hito", "Proyecto",
                                    "Responsable", "Estado", "Entregable"])
    h_first, h_rows = h_sec + 2, 30
    cond_hh = ('AND({a}<>"",OR({fd}="(Todos)",{d}={fd}),OR({fs}="(Todas)",{s}={fs}),'
               'OR({al}="Vencido",{al}="Esta semana",{al}="Próximos 30 días"))').format(
        a="Hitos!$%s{r}" % H_ID, d="Hitos!$%s{r}" % H_DEP, s="Hitos!$%s{r}" % H_SUB,
        al="Hitos!$%s{r}" % H_ALE, fd=FDEP, fs=FSUB)
    dynamic_list(ws, cond_hh, H_FIRST, H_LAST, "BR", "BS", "BP", h_first, h_rows)
    fill_dyn_table(ws, "BP", h_first, h_rows, [
        (H(H_FEC), FMT_DATE, "center"), (H(H_ALE), None, "center"),
        (H(H_NOM), None, "left"), (H(H_PROY), None, "left"), (H(H_RESP), None, "left"),
        (H(H_EST), None, "center"), (H(H_ENT), None, "left")])
    state_rules(ws, "C%d:C%d" % (h_first, h_first + h_rows - 1), ALERTA_COLORS)

    widths(ws, {"A": 3, "B": 15, "C": 36, "D": 20, "E": 16, "F": 12, "G": 13, "H": 12,
                "I": 13, "J": 11, "K": 11, "L": 12, "M": 12, "N": 14, "O": 13, "P": 11,
                "Q": 11, "R": 13})
    hide_helpers(ws, ["BH", "BJ", "BK", "BL", "BN", "BO", "BP", "BR", "BS"])
    ws.freeze_panes = "B12"
    return ws


# --------------------------------------------------------------------------
# Hoja Detalle proyecto (nivel proyecto)
# --------------------------------------------------------------------------
FICHA = [
    ("ID", P_ID, None), ("Departamento", P_DEP, None), ("Subárea", P_SUB, None),
    ("Owner", P_OWN, None), ("Tipo", P_TIPO, None), ("Prioridad", P_PRIO, None),
    ("Estado", P_EST, None), ("Fase", P_FASE, None),
    ("Fecha de inicio", P_INI, FMT_DATE), ("Fin previsto", P_FIN, FMT_DATE),
    ("Días restantes", P_DIAS, FMT_INT), ("Avance real", P_AVA, FMT_PCT0),
    ("Avance esperado por calendario", P_AVAE, FMT_PCT0),
    ("Desviación", P_DESV, '+0%;-0%;0%'), ("Semáforo", P_SEM, None),
    ("Riesgo", P_RIE, None), ("Presupuesto €", P_PPTO, FMT_EUR),
    ("Gasto €", P_GASTO, FMT_EUR), ("% presupuesto consumido", P_CONS, FMT_PCT0),
    ("% consecución media de KPIs", P_PKPI, FMT_PCT0),
    ("Objetivo de negocio", "P", None), ("Impacto esperado", "Q", None),
    ("Notas", "R", None),
]


def build_detalle(wb, dvs, proys):
    ws = wb.create_sheet("Detalle proyecto")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PURPLE_DARK

    title(ws, "B2", "Detalle de proyecto")
    subtitle(ws, "B3", "Elige un proyecto y verás su ficha completa, sus KPIs y sus hitos. "
                       "Es el último nivel de zoom: compañía → departamento → proyecto.")

    lc = ws.cell(row=5, column=2, value="Proyecto")
    lc.font = Font(name=FONT, size=9, bold=True, color=PURPLE)
    lc.alignment = Alignment(horizontal="right", vertical="center")
    sel = ws["C5"]
    sel.value = proys[0]["nombre"] if proys else ""
    sel.font = Font(name=FONT, size=12, bold=True, color=INK)
    sel.fill = PatternFill("solid", fgColor="FFF9DB")
    sel.border = thin(PURPLE)
    sel.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("C5:H5")
    ws.row_dimensions[5].height = 22
    add_dv(ws, "=%s" % P(P_NOM), "C5")

    ws["BD1"] = '=IFERROR(MATCH($C$5,%s,0),"")' % P(P_NOM)
    ws["BD2"] = '=IF($BD$1="","",INDEX(%s,$BD$1))' % P(P_ID)

    r = 7
    for label, col, fmt in FICHA:
        lab = ws.cell(row=r, column=2, value=label)
        lab.font = Font(name=FONT, size=9, bold=True, color=GREY)
        lab.alignment = Alignment(horizontal="right", vertical="top")
        val = ws.cell(row=r, column=3,
                      value='=IF($BD$1="","",INDEX(%s,$BD$1))' % P(col))
        val.font = Font(name=FONT, size=10, color=INK)
        val.alignment = Alignment(horizontal="left", vertical="top",
                                  wrap_text=col in ("P", "Q", "R"))
        val.border = Border(bottom=Side("thin", color=BORDER_C))
        if fmt:
            val.number_format = fmt
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 32 if col in ("P", "Q", "R") else 18
        r += 1
    lab = ws.cell(row=r, column=2, value="Hitos (total / completados / vencidos)")
    lab.font = Font(name=FONT, size=9, bold=True, color=GREY)
    lab.alignment = Alignment(horizontal="right", vertical="top")
    val = ws.cell(row=r, column=3, value=(
        '=IF($BD$1="","",INDEX({t},$BD$1)&"  /  "&INDEX({o},$BD$1)&"  /  "&INDEX({v},$BD$1))'
    ).format(t=P(P_NHIT), o=P(P_HOK), v=P(P_HVEN)))
    val.font = Font(name=FONT, size=10, color=INK)
    val.alignment = Alignment(horizontal="left", vertical="center")
    val.border = Border(bottom=Side("thin", color=BORDER_C))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)

    state_rules(ws, "C%d" % (7 + [f[1] for f in FICHA].index(P_EST)), ESTADO_COLORS)
    state_rules(ws, "C%d" % (7 + [f[1] for f in FICHA].index(P_SEM)), SEMAFORO_COLORS)
    state_rules(ws, "C%d" % (7 + [f[1] for f in FICHA].index(P_RIE)), RIESGO_COLORS)

    # --- KPIs del proyecto ---
    k_sec = r + 2
    section(ws, k_sec, 2, "KPIs del proyecto")
    table_header(ws, k_sec + 1, 2, ["ID", "KPI", "Unidad", "Dirección", "Baseline", "Target",
                                    "Actual", "% consecución", "Estado", "Frecuencia", "Fuente"])
    k_first, k_rows = k_sec + 2, 20
    cond_k = 'AND({a}<>"",{p}=$BD$2)'.format(a="KPIs!$%s{r}" % K_ID, p="KPIs!$%s{r}" % K_PID)
    dynamic_list(ws, cond_k, K_FIRST, K_LAST, "BJ", "BK", "BH", k_first, k_rows)
    fill_dyn_table(ws, "BH", k_first, k_rows, [
        (K(K_ID), None, "center"), (K(K_NOM), None, "left"), (K(K_UNI), None, "center"),
        (K(K_DIR), None, "center"), (K(K_BASE), FMT_NUM, "right"),
        (K(K_TGT), FMT_NUM, "right"), (K(K_ACT), FMT_NUM, "right"),
        (K(K_PCT), FMT_PCT0, "center"), (K(K_EST), None, "center"),
        (K(K_FREQ), None, "center"), (K(K_SRC), None, "center")])
    state_rules(ws, "J%d:J%d" % (k_first, k_first + k_rows - 1), KPI_COLORS)
    ws.conditional_formatting.add(
        "I%d:I%d" % (k_first, k_first + k_rows - 1),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=PURPLE))

    # --- Hitos del proyecto ---
    h_sec = k_first + k_rows + 1
    section(ws, h_sec, 2, "Hitos del proyecto")
    table_header(ws, h_sec + 1, 2, ["ID", "Hito", "Responsable", "Fecha objetivo", "Estado",
                                    "Días restantes", "Alerta", "Entregable"])
    h_first, h_rows = h_sec + 2, 30
    cond_h = 'AND({a}<>"",{p}=$BD$2)'.format(a="Hitos!$%s{r}" % H_ID, p="Hitos!$%s{r}" % H_PID)
    dynamic_list(ws, cond_h, H_FIRST, H_LAST, "BN", "BO", "BL", h_first, h_rows)
    fill_dyn_table(ws, "BL", h_first, h_rows, [
        (H(H_ID), None, "center"), (H(H_NOM), None, "left"), (H(H_RESP), None, "left"),
        (H(H_FEC), FMT_DATE, "center"), (H(H_EST), None, "center"),
        (H(H_DIAS), FMT_INT, "center"), (H(H_ALE), None, "center"),
        (H(H_ENT), None, "left")])
    state_rules(ws, "H%d:H%d" % (h_first, h_first + h_rows - 1), ALERTA_COLORS)

    widths(ws, {"A": 3, "B": 34, "C": 26, "D": 18, "E": 16, "F": 13, "G": 13, "H": 15,
                "I": 14, "J": 13, "K": 13, "L": 14})
    hide_helpers(ws, ["BD", "BH", "BJ", "BK", "BL", "BN", "BO"])
    ws.freeze_panes = "B6"
    return ws


# --------------------------------------------------------------------------
# Hoja Léeme
# --------------------------------------------------------------------------
LEEME = [
    ("h", "Growth OS — cómo funciona este fichero"),
    ("p", "Este Excel es la interfaz de trabajo de un sistema cuya fuente de la verdad vive "
          "en el repositorio de GitHub (growth-os/data/*.yaml). El fichero se genera desde "
          "esos datos y se puede volver a importar, de modo que cada estado del portfolio "
          "queda versionado y es auditable."),
    ("h2", "El ciclo de trabajo"),
    ("li", "1. Tú editas en Excel: estados, avances, gastos, fechas, KPIs, hitos. Sólo las "
           "columnas de fondo blanco."),
    ("li", "2. Guardas el fichero en su carpeta de OneDrive/SharePoint."),
    ("li", "3. Le dices al agente «sincroniza el Growth OS». El agente lee el Excel desde "
           "OneDrive, actualiza los YAML del repo y hace commit: ahí queda el histórico."),
    ("li", "4. Cuando hay cambios estructurales (nuevos departamentos, columnas, métricas "
           "o reglas de cálculo), el agente regenera el fichero y te lo entrega para que lo "
           "dejes en la carpeta. Sustituye el anterior y sigues trabajando."),
    ("h2", "Reglas de oro"),
    ("li", "No renombres las hojas ni las cabeceras de columna: el importador las busca por nombre."),
    ("li", "No escribas en las columnas marcadas con «ƒ» ni en las de fondo gris. Son fórmulas y "
           "se sobreescriben en cada regeneración."),
    ("li", "Para dar de baja una fila, vacíala o pon el estado en Cancelado. No borres la fila "
           "entera desde el medio de la tabla."),
    ("li", "El ID es la clave que une las tres hojas de datos. Si cambias un ID de proyecto, "
           "cambia también el «Proyecto ID» de sus KPIs e hitos."),
    ("li", "Los porcentajes se guardan como fracción: escribe 35% (o 0,35), no 35."),
    ("li", "Hay filas vacías de más en Proyectos, KPIs e Hitos para que puedas añadir registros "
           "sin regenerar nada: las fórmulas ya están puestas y se activan al rellenar el ID."),
    ("h2", "Qué hay en cada hoja"),
    ("li", "Dashboard — nivel compañía. Portfolio por estado, por departamento y por subárea, "
           "proyectos que requieren atención e hitos calientes."),
    ("li", "Zoom departamento — nivel departamento y subárea. Tres desplegables (amarillos) "
           "filtran las tarjetas y las tres tablas: proyectos, KPIs e hitos."),
    ("li", "Detalle proyecto — nivel proyecto. Ficha completa del proyecto elegido más sus "
           "KPIs y sus hitos."),
    ("li", "Proyectos / KPIs / Hitos — las tres tablas de entrada de datos. Aquí se escribe."),
    ("li", "Config — listas que alimentan los desplegables. Se regeneran desde taxonomia.yaml."),
    ("h2", "Cómo se calcula lo calculado"),
    ("li", "Avance esperado = (hoy − inicio) / (fin previsto − inicio), acotado entre 0% y 100%. "
           "Es el avance que tocaría tener si el proyecto fuera a ritmo de calendario."),
    ("li", "Desviación = avance real − avance esperado. Por debajo de −15pp el proyecto se marca "
           "como Desviado."),
    ("li", "Semáforo = Vencido si la fecha de fin ya pasó · Inminente si quedan 14 días o menos · "
           "Desviado si la desviación baja de −15pp · En plazo en el resto de casos."),
    ("li", "% consecución de un KPI = parte del recorrido baseline→target ya cubierta, teniendo en "
           "cuenta la dirección: si el KPI es de los que hay que bajar (CAC, CPL), el cálculo se "
           "invierte. 100% = target alcanzado."),
    ("li", "Estado de un KPI = Conseguido ≥100% · En camino ≥70% · Atención ≥30% · Crítico por debajo."),
    ("li", "Hito vencido = fecha objetivo anterior a hoy y estado distinto de Completado o Cancelado."),
    ("h2", "Añadir un departamento nuevo"),
    ("p", "El sistema está montado para escalar más allá de Growth. En el repo, añade el bloque "
          "del departamento y sus subáreas en growth-os/data/departamentos.yaml, y pídele al "
          "agente que regenere: el Dashboard, los desplegables y el Zoom lo recogen "
          "automáticamente, sin tocar fórmulas."),
    ("h2", "Datos de ejemplo"),
    ("p", "El fichero se entrega con 12 proyectos, 29 KPIs y 41 hitos de ejemplo: son realistas "
          "pero inventados, y están ahí para que se vea el sistema funcionando y para que sirvan "
          "de plantilla de formato. Sustitúyelos por los reales."),
]


def build_leeme(wb):
    ws = wb.create_sheet("Léeme")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = LAVENDER
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 118
    r = 2
    for kind, text in LEEME:
        c = ws.cell(row=r, column=2, value=text)
        if kind == "h":
            c.font = Font(name=FONT, size=20, bold=True, color=PURPLE_DARK)
            ws.row_dimensions[r].height = 30
        elif kind == "h2":
            r += 1
            c = ws.cell(row=r, column=2, value=text.upper())
            c.font = Font(name=FONT, size=10, bold=True, color=PURPLE)
            c.border = Border(bottom=Side("thin", color=PURPLE))
            ws.row_dimensions[r].height = 24
        elif kind == "li":
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            ws.row_dimensions[r].height = 15 * (1 + len(text) // 108)
        else:
            c.font = Font(name=FONT, size=10, color=GREY)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * (1 + len(text) // 108)
        r += 1
    # Leyenda de colores
    r += 1
    section(ws, r, 2, "Leyenda")
    r += 1
    for label, bg, fg in [("Celda editable (fondo blanco)", WHITE, INK),
                          ("Filtro editable", "FFF9DB", INK),
                          ("Columna calculada — no escribir (ƒ, fondo gris)", COMPUTED_BG, GREY),
                          ("Bien / conseguido / en plazo", GREEN_BG, GREEN_FG),
                          ("Atención / inminente / desviado", AMBER_BG, AMBER_FG),
                          ("Crítico / vencido / bloqueado", RED_BG, RED_FG)]:
        c = ws.cell(row=r, column=2, value="   " + label)
        c.font = Font(name=FONT, size=9, bold=True, color=fg)
        c.fill = PatternFill("solid", fgColor=bg)
        c.border = thin()
        r += 1
    return ws


# --------------------------------------------------------------------------
def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "dist" / "Founderz_Growth_OS.xlsx"
    tax, deps, proys, kpis, hitos = load()
    validate(tax, deps, proys, kpis, hitos)

    wb = Workbook()
    wb.remove(wb.active)
    dvs, _ = build_config(wb, tax, deps, proys)
    build_proyectos(wb, proys, dvs)
    build_kpis(wb, kpis, dvs)
    build_hitos(wb, hitos, dvs)
    build_dashboard(wb, tax, deps, dvs)
    build_zoom(wb, tax, dvs)
    build_detalle(wb, dvs, proys)
    build_leeme(wb)

    order = ["Proyectos", "KPIs", "Hitos", "Léeme", "Dashboard",
             "Zoom departamento", "Detalle proyecto", "Config"]
    wb._sheets = [wb[n] for n in order]
    wb.active = order.index("Dashboard")

    wb.properties.title = "Founderz Growth OS"
    wb.properties.creator = "Growth OS — growth-os/build_excel.py"
    wb.properties.description = ("Generado desde growth-os/data/*.yaml. No editar las columnas "
                                 "calculadas: se regeneran.")
    wb.calculation.fullCalcOnLoad = True

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("Escrito %s" % out)
    print("  %d proyectos · %d KPIs · %d hitos · %d departamentos" % (
        len(proys), len(kpis), len(hitos), len(deps)))


if __name__ == "__main__":
    main()
