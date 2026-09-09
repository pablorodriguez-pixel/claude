#!/usr/bin/env python3
"""
Growth OS — importador: trae de vuelta al repo lo que se haya editado en Excel.

    python3 growth-os/import_excel.py [entrada.xlsx] [--dry-run]

Lee las hojas Proyectos, KPIs e Hitos, valida, reescribe growth-os/data/*.yaml
conservando sus comentarios de cabecera, e imprime un resumen de los cambios
para poder contarlos en el commit.

Sólo se importan las columnas de entrada. Las columnas calculadas (ƒ) se
ignoran deliberadamente: se recalculan al regenerar el fichero.
"""
import datetime as dt
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"

# hoja -> (fichero yaml, [(cabecera en Excel, clave yaml)], clave de orden)
SPEC = {
    "Proyectos": ("proyectos.yaml", [
        ("ID", "id"), ("Proyecto", "nombre"), ("Departamento", "departamento"),
        ("Subárea", "subarea"), ("Owner", "owner"), ("Tipo", "tipo"),
        ("Prioridad", "prioridad"), ("Estado", "estado"), ("Fase", "fase"),
        ("Inicio", "fecha_inicio"), ("Fin previsto", "fecha_fin_prevista"),
        ("Avance %", "avance_pct"), ("Riesgo", "riesgo"),
        ("Presupuesto €", "presupuesto_eur"), ("Gasto €", "gasto_eur"),
        ("Objetivo de negocio", "objetivo"), ("Impacto esperado", "impacto_esperado"),
        ("Notas", "notas"),
    ]),
    "KPIs": ("kpis.yaml", [
        ("ID", "id"), ("Proyecto ID", "proyecto_id"), ("KPI", "nombre"),
        ("Unidad", "unidad"), ("Dirección", "direccion"), ("Baseline", "baseline"),
        ("Target", "target"), ("Actual", "actual"), ("Frecuencia", "frecuencia"),
        ("Fuente", "fuente"),
    ]),
    "Hitos": ("hitos.yaml", [
        ("ID", "id"), ("Proyecto ID", "proyecto_id"), ("Hito", "nombre"),
        ("Responsable", "responsable"), ("Fecha objetivo", "fecha_objetivo"),
        ("Estado", "estado"), ("Entregable", "entregable"),
    ]),
}
TEXT_KEYS = {"objetivo", "impacto_esperado", "notas", "entregable"}


def clean(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, float) and v.is_integer() and abs(v) >= 1:
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def read_sheet(ws, columns):
    header = {}
    for c in ws[1]:
        if isinstance(c.value, str) and c.value.strip():
            header.setdefault(c.value.strip(), c.column)
    missing = [h for h, _ in columns if h not in header]
    if missing:
        raise SystemExit("La hoja '%s' no tiene las columnas: %s.\n"
                         "No renombres las cabeceras generadas." % (ws.title, ", ".join(missing)))
    rows = []
    for r in range(2, ws.max_row + 1):
        rid = clean(ws.cell(row=r, column=header[columns[0][0]]).value)
        if rid in (None, ""):
            continue
        rec = {}
        for h, key in columns:
            v = clean(ws.cell(row=r, column=header[h]).value)
            if v is None:
                v = "" if key in TEXT_KEYS else None
            rec[key] = v
        rows.append(rec)
    return rows


def header_comment(path):
    """Conserva el bloque de comentarios del principio del YAML."""
    if not path.exists():
        return ""
    out = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("#") or not line.strip():
            out.append(line)
        else:
            break
    while out and not out[-1].strip():
        out.pop()
    return "\n".join(out) + "\n\n" if out else ""


def dump(path, rows, dry_run):
    body = yaml.dump(rows, allow_unicode=True, sort_keys=False,
                     default_flow_style=False, width=100)
    text = header_comment(path) + body
    if dry_run:
        return
    path.write_text(text, encoding="utf-8")


def diff(old, new, label):
    """Resumen legible de los cambios, para el mensaje de commit."""
    o = {r["id"]: r for r in old}
    n = {r["id"]: r for r in new}
    lines = []
    for i in sorted(set(n) - set(o)):
        lines.append("  + %s %s: %s" % (label, i, n[i].get("nombre", "")))
    for i in sorted(set(o) - set(n)):
        lines.append("  − %s %s: %s (eliminado)" % (label, i, o[i].get("nombre", "")))
    for i in sorted(set(o) & set(n)):
        for k in n[i]:
            a, b = o[i].get(k), n[i][k]
            if a != b:
                lines.append("  ~ %s %s · %s: %r → %r" % (label, i, k, a, b))
    return lines


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    src = Path(args[0]) if args else ROOT / "dist" / "Founderz_Growth_OS.xlsx"
    if not src.exists():
        raise SystemExit("No encuentro %s" % src)

    wb = load_workbook(src, data_only=True)
    changes, totals = [], []
    for sheet, (fname, columns) in SPEC.items():
        if sheet not in wb.sheetnames:
            raise SystemExit("El fichero no tiene la hoja '%s'." % sheet)
        rows = read_sheet(wb[sheet], columns)
        path = DATA / fname
        old = yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else []
        changes += diff(old or [], rows, sheet.rstrip("s"))
        totals.append("%d %s" % (len(rows), sheet.lower()))
        dump(path, rows, dry)

    print(("DRY RUN — " if dry else "") + "Importado desde %s" % src)
    print("  " + " · ".join(totals))
    if changes:
        print("\nCambios frente a los YAML del repo:")
        print("\n".join(changes[:200]))
        if len(changes) > 200:
            print("  ... y %d cambios más" % (len(changes) - 200))
    else:
        print("\nSin cambios frente a los YAML del repo.")
    print("\nSiguiente paso: python3 growth-os/build_excel.py  "
          "(para que el Excel recoja las columnas calculadas actualizadas)")


if __name__ == "__main__":
    main()
