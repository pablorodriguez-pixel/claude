#!/usr/bin/env python3
"""
Growth OS — variante «sólo datos» del libro.

    python3 growth-os/build_excel_datos.py [salida.xlsx]

Mismas tres tablas de entrada que el libro completo (Proyectos, KPIs, Hitos) con
sus desplegables y formatos, pero SIN fórmulas, sin columnas calculadas, sin
formato condicional y sin las hojas de dashboard.

Existe por una razón de transporte: un libro sin fórmulas pesa una fracción del
completo, y eso lo hace reemplazable por completo a través del conector de
Microsoft 365 — que sólo sabe de ficheros enteros. El análisis (semáforos,
desviaciones, % de consecución) vive en el dashboard web, que se genera de los
mismos datos.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_excel import (FMT_DATE, FMT_EUR, FMT_INT, FMT_NUM, FMT_PCT_IN, FONT,
                         PURPLE, UNIT_FMT, WHITE, add_dv, build_config, load,
                         thin, validate, widths)

SHEETS = {
    "Proyectos": (
        ["ID", "Proyecto", "Departamento", "Subárea", "Owner", "Tipo", "Prioridad",
         "Estado", "Fase", "Inicio", "Fin previsto", "Avance %", "Riesgo",
         "Presupuesto €", "Gasto €", "Objetivo de negocio", "Impacto esperado", "Notas"],
        ["id", "nombre", "departamento", "subarea", "owner", "tipo", "prioridad",
         "estado", "fase", "fecha_inicio", "fecha_fin_prevista", "avance_pct", "riesgo",
         "presupuesto_eur", "gasto_eur", "objetivo", "impacto_esperado", "notas"],
        {"A": 10, "B": 40, "C": 13, "D": 22, "E": 16, "F": 15, "G": 9, "H": 12,
         "I": 11, "J": 11, "K": 11, "L": 10, "M": 9, "N": 14, "O": 13,
         "P": 42, "Q": 42, "R": 42},
    ),
    "KPIs": (
        ["ID", "Proyecto ID", "KPI", "Unidad", "Dirección", "Baseline", "Target",
         "Actual", "Frecuencia", "Fuente"],
        ["id", "proyecto_id", "nombre", "unidad", "direccion", "baseline", "target",
         "actual", "frecuencia", "fuente"],
        {"A": 10, "B": 12, "C": 34, "D": 10, "E": 11, "F": 12, "G": 12, "H": 12,
         "I": 12, "J": 15},
    ),
    "Hitos": (
        ["ID", "Proyecto ID", "Hito", "Responsable", "Fecha objetivo", "Estado", "Entregable"],
        ["id", "proyecto_id", "nombre", "responsable", "fecha_objetivo", "estado", "entregable"],
        {"A": 9, "B": 12, "C": 44, "D": 16, "E": 14, "F": 13, "G": 40},
    ),
}
SLACK = {"Proyectos": 28, "KPIs": 40, "Hitos": 40}
FECHAS = {"fecha_inicio", "fecha_fin_prevista", "fecha_objetivo"}
EUROS = {"presupuesto_eur", "gasto_eur"}
ANCHO = {"objetivo", "impacto_esperado", "notas", "entregable"}


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "dist" / "Growth_OS_Datos.xlsx"
    tax, deps, proys, kpis, hitos = load()
    validate(tax, deps, proys, kpis, hitos)
    filas = {"Proyectos": proys, "KPIs": kpis, "Hitos": hitos}

    wb = Workbook()
    wb.remove(wb.active)
    dvs, _ = build_config(wb, tax, deps, proys)

    for name, (heads, keys, w) in SHEETS.items():
        ws = wb.create_sheet(name)
        for i, h in enumerate(heads, start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = Font(name=FONT, size=9, bold=True, color=WHITE)
            c.fill = PatternFill("solid", fgColor=PURPLE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        for r, rec in enumerate(filas[name], start=2):
            for i, key in enumerate(keys, start=1):
                c = ws.cell(row=r, column=i, value=rec.get(key, ""))
                c.font = Font(name=FONT, size=9)
                c.border = thin()
                c.alignment = Alignment(vertical="center", wrap_text=key in ANCHO,
                                        horizontal="left" if key in ANCHO or key == "nombre" else "center")
                if key in FECHAS:
                    c.number_format = FMT_DATE
                elif key in EUROS:
                    c.number_format = FMT_EUR
                elif key == "avance_pct":
                    c.number_format = FMT_PCT_IN
                elif key in ("baseline", "target", "actual"):
                    c.number_format = UNIT_FMT.get(rec.get("unidad"), FMT_NUM)

        last = len(filas[name]) + 1 + SLACK[name]
        widths(ws, w)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(heads)), last)
        ws.sheet_properties.tabColor = PURPLE

        col = {k: get_column_letter(i) for i, k in enumerate(keys, start=1)}
        rng = lambda k: "%s2:%s%d" % (col[k], col[k], last)
        if name == "Proyectos":
            for k, lista in (("departamento", "Departamento"), ("subarea", "Subárea"),
                             ("owner", "Owner"), ("tipo", "Tipo"), ("prioridad", "Prioridad"),
                             ("estado", "Estado proyecto"), ("fase", "Fase"), ("riesgo", "Riesgo")):
                add_dv(ws, dvs[lista], rng(k))
        elif name == "KPIs":
            add_dv(ws, "=Proyectos!$A$2:$A$%d" % (len(proys) + 29), rng("proyecto_id"))
            for k, lista in (("unidad", "Unidad KPI"), ("direccion", "Dirección KPI"),
                             ("frecuencia", "Frecuencia KPI"), ("fuente", "Fuente dato")):
                add_dv(ws, dvs[lista], rng(k))
        else:
            add_dv(ws, "=Proyectos!$A$2:$A$%d" % (len(proys) + 29), rng("proyecto_id"))
            add_dv(ws, dvs["Owner"], rng("responsable"))
            add_dv(ws, dvs["Estado hito"], rng("estado"))

    wb._sheets = [wb[n] for n in ["Proyectos", "KPIs", "Hitos", "Config"]]
    wb.active = 0
    wb.properties.title = "Founderz Growth OS — datos"
    wb.properties.description = ("Tablas de entrada. El análisis vive en el dashboard web, "
                                "generado de estos mismos datos.")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("Escrito %s" % out)


if __name__ == "__main__":
    main()
