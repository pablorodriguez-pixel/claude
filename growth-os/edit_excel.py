#!/usr/bin/env python3
"""
Growth OS — ediciones puntuales sobre el libro, sin regenerarlo.

    python3 growth-os/edit_excel.py LIBRO.xlsx --ops '[{"id":"GRW-004","estado":"En curso"}]'
    python3 growth-os/edit_excel.py LIBRO.xlsx --ops-file cambios.json [--dry-run]

Esta es la herramienta de escritura del agente. A diferencia de build_excel.py,
que reconstruye el fichero desde los YAML, esto abre el libro que el equipo está
usando y toca sólo las celdas indicadas. Todo lo demás —filas añadidas a mano,
notas, formatos -- se queda como está.

Reglas que impone:
  · Sólo escribe en columnas de entrada. Las columnas ƒ son fórmulas y las
    rechaza, porque sobreescribirlas rompería el recálculo.
  · Valida contra taxonomia.yaml: un estado que no existe se rechaza antes
    de tocar el fichero.
  · Nada se escribe si alguna operación es inválida. O todas o ninguna.

Formato de una operación:
  {"id": "GRW-004", "estado": "En curso", "avance_pct": 0.7}
  {"id": "H-015", "estado": "Completado"}
  {"nuevo": "proyecto", "id": "GRW-013", "nombre": "...", ...}   ← añade una fila

El id determina la hoja: se busca en Proyectos, KPIs e Hitos por la columna ID.
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"

sys.path.insert(0, str(ROOT))
from import_excel import SPEC, TEXT_KEYS  # reutiliza el mapa cabecera → clave

# Qué campo de taxonomia.yaml valida cada clave editable
VALIDA = {
    "estado": {"Proyectos": "estado_proyecto", "Hitos": "estado_hito"},
    "fase": {"Proyectos": "fase"},
    "prioridad": {"Proyectos": "prioridad"},
    "tipo": {"Proyectos": "tipo_proyecto"},
    "riesgo": {"Proyectos": "riesgo"},
    "unidad": {"KPIs": "unidad_kpi"},
    "direccion": {"KPIs": "direccion_kpi"},
    "frecuencia": {"KPIs": "frecuencia_kpi"},
    "fuente": {"KPIs": "fuente_dato"},
}
FECHAS = {"fecha_inicio", "fecha_fin_prevista", "fecha_objetivo"}
FRACCIONES = {"avance_pct"}
HOJA_DE = {"proyecto": "Proyectos", "kpi": "KPIs", "hito": "Hitos"}


def coerce(key, v):
    """Normaliza el valor a lo que Excel espera en esa columna."""
    if v is None or v == "":
        return None
    if key in FECHAS:
        if isinstance(v, (dt.date, dt.datetime)):
            return v if isinstance(v, dt.date) else v.date()
        return dt.date.fromisoformat(str(v).strip())
    if key in FRACCIONES:
        s = str(v).strip()
        if s.endswith("%"):
            return float(s[:-1].replace(",", ".")) / 100
        f = float(s)
        return f / 100 if f > 1 else f      # acepta 70 y 0.7 como lo mismo
    if key.endswith("_eur") or key in {"baseline", "target", "actual"}:
        s = str(v).replace("€", "").replace(".", "").replace(",", ".").strip()
        return float(s) if "." in s else int(s or 0)
    return str(v)


def index_sheet(ws, columns):
    """Devuelve (cabecera → columna, id → fila, primera fila libre)."""
    header = {}
    for c in ws[1]:
        if isinstance(c.value, str) and c.value.strip():
            header.setdefault(c.value.strip(), c.column)
    missing = [h for h, _ in columns if h not in header]
    if missing:
        raise SystemExit("La hoja '%s' no tiene las columnas: %s" % (ws.title, ", ".join(missing)))
    idcol = header[columns[0][0]]
    rows, libre = {}, None
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=idcol).value
        if v in (None, ""):
            if libre is None:
                libre = r
        else:
            rows[str(v).strip()] = r
    return header, rows, libre if libre else ws.max_row + 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("libro")
    ap.add_argument("--ops", help="JSON con la lista de operaciones")
    ap.add_argument("--ops-file", help="fichero JSON con la lista de operaciones")
    ap.add_argument("--dry-run", action="store_true", help="enseña los cambios y no escribe")
    a = ap.parse_args()

    if not a.ops and not a.ops_file:
        raise SystemExit("Hace falta --ops o --ops-file")
    ops = json.loads(a.ops) if a.ops else json.loads(Path(a.ops_file).read_text(encoding="utf-8"))
    if isinstance(ops, dict):
        ops = [ops]

    tax = yaml.safe_load((DATA / "taxonomia.yaml").read_text(encoding="utf-8"))
    path = Path(a.libro)
    if not path.exists():
        raise SystemExit("No encuentro %s" % path)
    wb = load_workbook(path)          # data_only=False: conserva las fórmulas

    idx = {}
    for sheet, (_, columns) in SPEC.items():
        if sheet not in wb.sheetnames:
            raise SystemExit("El libro no tiene la hoja '%s'." % sheet)
        header, rows, libre = index_sheet(wb[sheet], columns)
        idx[sheet] = {"ws": wb[sheet], "header": header, "rows": rows, "libre": libre,
                      "keys": {k: h for h, k in columns}}

    plan, errores = [], []
    for n, op in enumerate(ops, 1):
        rid = str(op.get("id", "")).strip()
        if not rid:
            errores.append("op %d: falta el id" % n)
            continue
        nuevo = op.get("nuevo")
        if nuevo:
            sheet = HOJA_DE.get(str(nuevo).lower())
            if not sheet:
                errores.append("op %d: 'nuevo' debe ser proyecto, kpi o hito" % n)
                continue
            if rid in idx[sheet]["rows"]:
                errores.append("op %d: el id %s ya existe en %s" % (n, rid, sheet))
                continue
            fila = idx[sheet]["libre"]
            idx[sheet]["libre"] += 1
            idx[sheet]["rows"][rid] = fila
        else:
            hojas = [s for s in idx if rid in idx[s]["rows"]]
            if not hojas:
                errores.append("op %d: no encuentro el id %s en ninguna hoja" % (n, rid))
                continue
            sheet, fila = hojas[0], idx[hojas[0]]["rows"][rid]

        for key, raw in op.items():
            if key == "nuevo":
                continue
            if key == "id" and not nuevo:
                continue          # reescribir el id con su mismo valor es ruido
            head = idx[sheet]["keys"].get(key)
            if not head:
                errores.append("op %d: '%s' no es una columna editable de %s" % (n, key, sheet))
                continue
            if key in VALIDA and sheet in VALIDA[key]:
                permitidos = tax[VALIDA[key][sheet]]
                if raw not in permitidos:
                    errores.append("op %d: %s='%s' no está en la taxonomía (%s)" % (
                        n, key, raw, ", ".join(permitidos)))
                    continue
            try:
                val = coerce(key, raw)
            except Exception as e:
                errores.append("op %d: no puedo interpretar %s='%s' (%s)" % (n, key, raw, e))
                continue
            col = idx[sheet]["header"][head]
            antes = idx[sheet]["ws"].cell(row=fila, column=col).value
            plan.append((sheet, fila, col, head, rid, antes, val))

    if errores:
        print("No he escrito nada. Errores:", file=sys.stderr)
        for e in errores:
            print("  - " + e, file=sys.stderr)
        sys.exit(1)

    if not plan:
        print("Nada que cambiar.")
        return

    print(("DRY RUN — " if a.dry_run else "") + "%d celda(s):" % len(plan))
    for sheet, fila, col, head, rid, antes, val in plan:
        print("  %s · %s · fila %d · %s: %r → %r" % (sheet, rid, fila, head, antes, val))

    if a.dry_run:
        return
    for sheet, fila, col, head, rid, antes, val in plan:
        idx[sheet]["ws"].cell(row=fila, column=col, value=val)
    wb.save(path)
    print("\nGuardado %s" % path)
    print("Las columnas ƒ se recalculan al abrir el libro en Excel.")


if __name__ == "__main__":
    main()
